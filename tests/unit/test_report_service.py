"""F6: report export service — ownership, validation, real xlsx output, row cap."""

import sqlite3

import pytest
from unittest.mock import patch

from app.models.chat import ChatHistory
from app.models.report_export import ReportExport
from app.services import report_service
from app.services.report_service import ExportError, cleanup_expired, create_export, run_export


@pytest.fixture
def chat_with_sql(db_session, test_user):
    chat = ChatHistory(
        user_id=test_user.id,
        question="รายได้รวม",
        generated_sql="SELECT * FROM revenue",
        ai_response="...",
    )
    db_session.add(chat)
    db_session.commit()
    db_session.refresh(chat)
    return chat


@pytest.fixture
def business_db(tmp_path, monkeypatch):
    """Temp business sqlite with 20 rows; export dir redirected to tmp."""
    db_path = tmp_path / "biz.sqlite"
    conn = sqlite3.connect(db_path)
    conn.execute('CREATE TABLE revenue (id INTEGER, "กลุ่มธุรกิจ" TEXT, val REAL)')
    conn.executemany("INSERT INTO revenue VALUES (?, ?, ?)", [(i, f"กลุ่ม{i}", i * 1.5) for i in range(20)])
    conn.commit()
    conn.close()

    from app.config import settings
    monkeypatch.setattr(settings, "BUSINESS_DB_PATH", str(db_path))
    monkeypatch.setattr(report_service, "EXPORT_DIR", tmp_path / "exports")
    return db_path


class TestCreateExport:
    def test_ownership_enforced(self, db_session, chat_with_sql, admin_user):
        other = type("U", (), {"id": 9999, "role": "user"})()
        with pytest.raises(ExportError, match="สิทธิ์"):
            create_export(db_session, other, chat_with_sql.id)
        # Admin can export anyone's chat
        export = create_export(db_session, admin_user, chat_with_sql.id)
        assert export.status == "pending"

    def test_dangerous_sql_rejected(self, db_session, test_user):
        chat = ChatHistory(user_id=test_user.id, question="x", generated_sql="DROP TABLE revenue", ai_response="")
        db_session.add(chat)
        db_session.commit()
        with pytest.raises(ExportError, match="ไม่ผ่านการตรวจสอบ"):
            create_export(db_session, test_user, chat.id)

    def test_missing_sql_rejected(self, db_session, test_user):
        chat = ChatHistory(user_id=test_user.id, question="x", generated_sql=None, ai_response="")
        db_session.add(chat)
        db_session.commit()
        with pytest.raises(ExportError, match="ไม่มี SQL"):
            create_export(db_session, test_user, chat.id)


class TestRunExport:
    def test_export_produces_valid_xlsx(self, db_session, test_user, chat_with_sql, business_db):
        export = create_export(db_session, test_user, chat_with_sql.id)
        with patch.object(report_service, "_config_int", side_effect=lambda k, d: d):
            run_export(export.id, db=db_session)
        db_session.refresh(export)
        assert export.status == "done", export.error
        assert export.row_count == 20
        assert export.truncated is False

        from openpyxl import load_workbook
        wb = load_workbook(export.file_path)
        assert set(wb.sheetnames) == {"Data", "Info"}
        data = wb["Data"]
        assert data.max_row == 21  # header + 20 rows
        assert data.cell(1, 2).value == "กลุ่มธุรกิจ"  # Thai header intact

    def test_row_cap_truncates(self, db_session, test_user, chat_with_sql, business_db):
        export = create_export(db_session, test_user, chat_with_sql.id)
        with patch.object(report_service, "_config_int",
                          side_effect=lambda k, d: 5 if k == "export_max_rows" else d):
            run_export(export.id, db=db_session)
        db_session.refresh(export)
        assert export.status == "done"
        assert export.row_count == 5
        assert export.truncated is True

    def test_failure_sets_failed_status(self, db_session, test_user, business_db):
        chat = ChatHistory(user_id=test_user.id, question="x",
                           generated_sql="SELECT * FROM no_such_table", ai_response="")
        db_session.add(chat)
        db_session.commit()
        export = create_export(db_session, test_user, chat.id)
        run_export(export.id, db=db_session)
        db_session.refresh(export)
        assert export.status == "failed"
        assert export.error


class TestCleanup:
    def test_expired_record_and_file_removed(self, db_session, test_user, chat_with_sql, business_db):
        from datetime import timedelta
        from app.core.time_utils import utcnow

        export = create_export(db_session, test_user, chat_with_sql.id)
        with patch.object(report_service, "_config_int", side_effect=lambda k, d: d):
            run_export(export.id, db=db_session)
        db_session.refresh(export)
        file_path = export.file_path

        export.expires_at = utcnow() - timedelta(days=1)
        db_session.commit()

        removed = cleanup_expired(db_session)
        assert removed >= 1
        import os
        assert not os.path.exists(file_path)
        assert db_session.query(ReportExport).filter(ReportExport.id == export.id).first() is None
